#!/usr/bin/env python3
"""Extract the supplied roadmap without inventing evidence, ratings, or equations.

Requires python-docx only for re-extraction. All table and row locators are 1-based.
"""
from pathlib import Path
import argparse, hashlib, json, re
from docx import Document
from docx.oxml.ns import qn

GROUPS = [
 ('A', 'Reference-rich foundations and early commercial candidates', 21, (13, 14)),
 ('B', 'Exact anchors plus manageable numerical qualification', 22, (14, 16)),
 ('C', 'General engineering regimes requiring stronger numerical and physical evidence', 23, (16, 17)),
 ('D', 'Hard regimes where model form, topology, chaos or experiment dominates', 24, (17, 19)),
]

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()

def extract(path):
 doc=Document(path)
 rows=[]
 for group, name, ti, pages in GROUPS:
  t=doc.tables[ti]
  assert len(t.columns)==5 and t.rows[0].cells[0].text=='Physics regime and engineering jobs'
  for ri, row in enumerate(t.rows[1:], 2):
   raw=[c.text for c in row.cells]
   title, sep, jobs=raw[0].partition('\n')
   match=re.fullmatch(r'(Q[1-5]) \| Cost ([1-5]) \| Hidden ([1-5]) \| GTM ([1-5]\*?)',raw[3])
   if not match: raise ValueError(f'Unexpected source rating at table {ti+1} row {ri}: {raw[3]}')
   rows.append({
    'id':f'PHY-{group}{ri-1:02d}', 'group':group,'group_title':name,
    'title':title, 'engineering_jobs_source':jobs,
    'reference_path_source':raw[1], 'mms_hybrid_source':raw[2],
    'ratings_source':raw[3], 'path_and_risk_source':raw[4],
    'difficulty_hypothesis':match[1], 'mms_rating_source':re.match(r'MMS ([^.]+)\.', raw[2])[1],
    'source_locator':{'section':'8','docx_table_1based':ti+1,'row_1based':ri,'rendered_page_range':list(pages)},
    'source_cells':raw,
    'epistemic_status':'SOURCE_RESEARCH_HYPOTHESIS',
    'profile_status':'NOT_PROFILED', 'qualified_carbon_evidence':[],
    'equations_to_register':None,'boundary_conditions_to_register':None,
    'physical_envelope_to_register':None,'measured_operating_profile':None,
   })
 def table(ti):
  t=doc.tables[ti]; keys=[c.text for c in t.rows[0].cells]
  return [{'source_locator':{'docx_table_1based':ti+1,'row_1based':i}, **dict(zip(keys,[c.text for c in row.cells]))} for i,row in enumerate(t.rows[1:],2)]
 refs=[]
 for p in doc.paragraphs:
  m=re.match(r'^\[([CE]\d+)\]',p.text)
  if not m: continue
  urls=[]
  for h in p._p.findall(qn('w:hyperlink')):
   rid=h.get(qn('r:id'))
   if rid and rid in doc.part.rels: urls.append(doc.part.rels[rid].target_ref)
  refs.append({'id':m[1],'citation_source':p.text,'urls_in_source':urls,'verification_status':'AS_CITED_NOT_REVERIFIED'})
 return {
  'schema_version':'carbon_opportunity_atlas_source_v0.1',
  'source':{'id':'GTM-2026-08-27','file':path.name,'sha256':sha(path),
     'research_date':'2026-08-27','source_repo_snapshot':'4e4a66d29566a2a62a82188adddac76e6e0fb8b8',
     'status':'Working research memo, not scientific qualification, customer demand, or implementation authority.'},
  'extraction':{'method':'Exact python-docx table-cell text; title/jobs split at first newline; ratings parsed without new ranking.',
     'opportunities':len(rows),'page_locator_note':'Page ranges refer to the supplied rendered document. Table/row are stable extraction locators.',
     'source_recommendation_note':'BUILD NOW, PILOT, and Q/Cost/Hidden/GTM labels remain historical research hypotheses, not active ticket selection.'},
  'groups':[{'id':g,'title':n} for g,n,_,_ in GROUPS],
  'opportunities':rows,
  'reference_families':table(8), 'mms_roles':table(12),'mms_factory':table(13),
  'mms_limits':table(11),'product_archetypes':table(5),'job_archetypes':table(18),
  'portfolio_tracks':table(19),'qualification_stages':table(25),
  'experiments':table(28),'thermal_ladder':table(31),'bibliography':refs,
 }

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('docx',type=Path);p.add_argument('output',type=Path);args=p.parse_args()
 data=extract(args.docx); args.output.parent.mkdir(parents=True,exist_ok=True)
 args.output.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 print(f"Extracted {len(data['opportunities'])} opportunities, {len(data['reference_families'])} reference families, {len(data['experiments'])} proposed experiments.")
